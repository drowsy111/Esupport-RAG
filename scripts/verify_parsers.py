"""Parser checks using existing WSL dependencies; never installs packages or reads user files."""
import io
import unittest

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.core.exceptions import AppError
from app.ingestion.parsers.registry import parse_document


def text_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (DEMO-DB-001 MySQL 8.0) Tj ET")
    page[NameObject("/Contents")] = stream
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class ParserChecks(unittest.TestCase):
    def test_pdf_page_and_exact_identifier(self):
        parsed = parse_document("pdf", text_pdf())
        self.assertIn("DEMO-DB-001", parsed.blocks[0].text)
        self.assertEqual(parsed.blocks[0].page, 1)
        self.assertEqual(parsed.blocks[0].anchor, "page-1")

    def test_scanned_or_blank_pdf_reports_ocr(self):
        writer, output = PdfWriter(), io.BytesIO()
        writer.add_blank_page(width=300, height=300)
        writer.write(output)
        with self.assertRaises(AppError) as result:
            parse_document("pdf", output.getvalue())
        self.assertEqual(result.exception.code, "OCR_REQUIRED")

    def test_docx_preserves_order_and_uses_block_not_fake_page(self):
        document = Document()
        document.add_heading("Database setup", level=1)
        document.add_paragraph("Before table")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "DB_HOST"
        table.cell(0, 1).text = "mysql"
        document.add_paragraph("After table")
        output = io.BytesIO()
        document.save(output)
        parsed = parse_document("docx", output.getvalue())
        self.assertEqual([b.text for b in parsed.blocks][1:],
                         ["Before table", "DB_HOST | mysql", "After table"])
        self.assertTrue(all(b.page is None for b in parsed.blocks))
        self.assertEqual(parsed.blocks[-1].anchor, "block-4")

    def test_markdown_code_is_not_interpreted_as_heading(self):
        fence = chr(96) * 3
        content = f"# Real heading\n\n{fence}sh\n# comment\n{fence}\n"
        parsed = parse_document("md", content.encode())
        self.assertTrue(all(b.title_path == "Real heading" for b in parsed.blocks))


if __name__ == "__main__":
    unittest.main(verbosity=2)

